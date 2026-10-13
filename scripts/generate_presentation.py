from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Brand colors
OFF_WHITE = RGBColor(0xFA, 0xFA, 0xFA)
CORAL = RGBColor(0xE0, 0x7B, 0x5A)
CORAL_LIGHT = RGBColor(0xF2, 0xC4, 0xB3)
CHARCOAL = RGBColor(0x2D, 0x2D, 0x2D)
SLATE = RGBColor(0x6B, 0x7A, 0x8D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xE8, 0xE8, 0xE8)
DARK_CORAL = RGBColor(0xC4, 0x5F, 0x3A)
PALE_CORAL = RGBColor(0xFB, 0xEE, 0xE8)
BLUE = RGBColor(0x3A, 0x7B, 0xD5)
PALE_BLUE = RGBColor(0xE8, 0xF1, 0xFC)
DARK_BLUE = RGBColor(0x1E, 0x55, 0xAA)
GREEN = RGBColor(0x4A, 0x9A, 0x6A)

TOTAL_SLIDES = 16

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

HEADING_FONT = "Georgia"
BODY_FONT = "Calibri"


def add_rect(slide, left, top, width, height, fill_color=None, line_color=None, line_width=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        if line_width:
            shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape


def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, text, left, top, width, height,
                 font_size=14, bold=False, color=None, align=PP_ALIGN.LEFT,
                 font_name=None, italic=False, wrap=True):
    if font_name is None:
        font_name = BODY_FONT
    txBox = slide.shapes.add_textbox(left, top, width, height)
    txBox.word_wrap = wrap
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font_name
    if color:
        run.font.color.rgb = color
    return txBox


def add_header_bar(slide, title, accent_color=CORAL):
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), fill_color=accent_color)
    add_rect(slide, 0, Inches(0.08), SLIDE_W, Inches(0.85), fill_color=CHARCOAL)
    add_text_box(slide, title,
                 Inches(0.4), Inches(0.1), Inches(10), Inches(0.9),
                 font_size=22, bold=True, color=WHITE, font_name=HEADING_FONT)
    add_text_box(slide, "SYNERGASIA",
                 Inches(10.8), Inches(0.1), Inches(2.3), Inches(0.9),
                 font_size=11, bold=True, color=CORAL, align=PP_ALIGN.RIGHT, font_name=HEADING_FONT)


def add_footer(slide, page_num):
    add_rect(slide, 0, Inches(7.15), SLIDE_W, Inches(0.35), fill_color=LIGHT_GRAY)
    add_text_box(slide, "Synergasia — L0 Platform Overview  |  Confidential",
                 Inches(0.3), Inches(7.17), Inches(9), Inches(0.3),
                 font_size=8, color=SLATE)
    add_text_box(slide, f"{page_num} / {TOTAL_SLIDES}",
                 Inches(12.3), Inches(7.17), Inches(0.8), Inches(0.3),
                 font_size=8, color=SLATE, align=PP_ALIGN.RIGHT)


def add_section_label(slide, label, left, top, width, color=CORAL):
    add_rect(slide, left, top, width, Inches(0.32), fill_color=color)
    add_text_box(slide, label,
                 left + Inches(0.12), top + Inches(0.04), width - Inches(0.24), Inches(0.28),
                 font_size=10, bold=True, color=WHITE)


blank_layout = prs.slide_layouts[6]


# ─────────────────────────────────────────
# SLIDE 1: COVER
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)

add_rect(slide, 0, 0, Inches(4.8), SLIDE_H, fill_color=CORAL)

circle = slide.shapes.add_shape(9, Inches(2.8), Inches(4.5), Inches(3.5), Inches(3.5))
circle.fill.solid()
circle.fill.fore_color.rgb = DARK_CORAL
circle.line.fill.background()

add_rect(slide, Inches(4.8), Inches(5.5), Inches(8.53), Inches(0.06), fill_color=CORAL_LIGHT)

add_text_box(slide, "Synergasia",
             Inches(0.35), Inches(1.5), Inches(4.1), Inches(1.5),
             font_size=46, bold=True, color=WHITE, font_name=HEADING_FONT)
add_text_box(slide, "Guiding students toward their future",
             Inches(0.35), Inches(3.0), Inches(4.1), Inches(0.8),
             font_size=14, color=WHITE, italic=True, font_name=HEADING_FONT)

add_text_box(slide, "Platform Overview",
             Inches(5.3), Inches(1.8), Inches(7.5), Inches(0.7),
             font_size=32, bold=True, color=CHARCOAL, font_name=HEADING_FONT)
add_text_box(slide, "L0 — Stakeholder Briefing",
             Inches(5.3), Inches(2.55), Inches(7.5), Inches(0.5),
             font_size=18, color=SLATE, font_name=HEADING_FONT)

add_rect(slide, Inches(5.3), Inches(3.15), Inches(1.2), Inches(0.05), fill_color=CORAL)

add_text_box(slide, "A comprehensive student counseling platform with\nrole-based access for Admins, Counselors, and Students.",
             Inches(5.3), Inches(3.35), Inches(7.5), Inches(1.0),
             font_size=13, color=CHARCOAL)
add_text_box(slide, "June 2026",
             Inches(5.3), Inches(6.4), Inches(3), Inches(0.4),
             font_size=11, color=SLATE)

add_footer(slide, 1)


# ─────────────────────────────────────────
# SLIDE 2: PLATFORM OVERVIEW
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "What is Synergasia?")
add_footer(slide, 2)

personas = [
    ("Admin", CORAL, [
        "Manages counselors & students",
        "Configures platform settings",
        "AI-powered data extraction",
        "Full oversight & audit trail",
    ]),
    ("Counselor", SLATE, [
        "Views assigned students only",
        "Assigns tasks & gives feedback",
        "Uploads documents securely",
        "Coordinates via planner",
    ]),
    ("Student", BLUE, [
        "Tracks profile completion",
        "Manages activities & portfolio",
        "Plans university applications",
        "Receives push notifications",
    ]),
]

card_width = Inches(3.8)
for i, (name, color, bullets) in enumerate(personas):
    left = Inches(0.4) + i * Inches(4.05)
    add_rect(slide, left, Inches(1.1), card_width, Inches(5.6), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_rect(slide, left, Inches(1.1), card_width, Inches(0.5), fill_color=color)
    add_text_box(slide, name,
                 left + Inches(0.15), Inches(1.12), card_width - Inches(0.3), Inches(0.45),
                 font_size=16, bold=True, color=WHITE, font_name=HEADING_FONT)
    for j, b in enumerate(bullets):
        add_text_box(slide, f"▸  {b}",
                     left + Inches(0.2), Inches(1.8) + j * Inches(0.7),
                     card_width - Inches(0.3), Inches(0.6),
                     font_size=12, color=CHARCOAL)

add_rect(slide, Inches(0.4), Inches(6.55), Inches(12.5), Inches(0.55), fill_color=PALE_CORAL)
highlights = ["Multi-tenant architecture", "Role-based access control (RBAC)", "AI-powered features", "Real-time push notifications"]
for i, h in enumerate(highlights):
    add_text_box(slide, f"✓  {h}",
                 Inches(0.6) + i * Inches(3.1), Inches(6.6), Inches(2.9), Inches(0.4),
                 font_size=10, bold=True, color=DARK_CORAL)


# ─────────────────────────────────────────
# SLIDE 3: ADMIN SECTION DIVIDER
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, CHARCOAL)

add_rect(slide, 0, 0, Inches(0.5), SLIDE_H, fill_color=CORAL)
add_rect(slide, Inches(0.5), 0, Inches(5.5), SLIDE_H, fill_color=DARK_CORAL)

circle2 = slide.shapes.add_shape(9, Inches(4.5), Inches(-1.0), Inches(5.5), Inches(5.5))
circle2.fill.solid()
circle2.fill.fore_color.rgb = RGBColor(0xB8, 0x52, 0x30)
circle2.line.fill.background()

add_text_box(slide, "01",
             Inches(0.8), Inches(1.5), Inches(2.0), Inches(2.0),
             font_size=96, bold=True, color=WHITE, font_name=HEADING_FONT)
add_text_box(slide, "ADMIN",
             Inches(0.8), Inches(3.6), Inches(4.5), Inches(0.8),
             font_size=36, bold=True, color=WHITE, font_name=HEADING_FONT)
add_rect(slide, Inches(0.8), Inches(4.5), Inches(1.5), Inches(0.06), fill_color=WHITE)
add_text_box(slide, "Command, customise, and\nintelligently oversee the\nentire platform.",
             Inches(0.8), Inches(4.7), Inches(4.5), Inches(1.5),
             font_size=15, color=CORAL_LIGHT, font_name=HEADING_FONT, italic=True)

add_text_box(slide, "3 slides in this section:",
             Inches(6.5), Inches(2.2), Inches(6.5), Inches(0.4),
             font_size=12, color=SLATE, bold=True)
section_slides = [
    "Command & Control — manage counselors, students, and access",
    "Customise the Platform — dynamic fields, university repository",
    "Intelligence & Oversight — AI extraction, email, notifications",
]
for i, s in enumerate(section_slides):
    add_rect(slide, Inches(6.5), Inches(2.75) + i * Inches(1.2), Inches(6.3), Inches(1.0),
             fill_color=RGBColor(0x3A, 0x3A, 0x3A), line_color=CORAL, line_width=Pt(0.5))
    add_text_box(slide, f"{i+1}.  {s}",
                 Inches(6.7), Inches(2.85) + i * Inches(1.2), Inches(5.9), Inches(0.8),
                 font_size=12, color=WHITE)

add_footer(slide, 3)


# ─────────────────────────────────────────
# SLIDE 4: ADMIN — Command & Control
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Admin: Command & Control")
add_footer(slide, 4)

add_section_label(slide, "ADMIN  ·  1 of 3", Inches(0.4), Inches(1.05), Inches(2.2))

add_text_box(slide, "Full platform control from a single command center.",
             Inches(0.4), Inches(1.5), Inches(5.8), Inches(0.9),
             font_size=15, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

txBox = slide.shapes.add_textbox(Inches(0.4), Inches(2.5), Inches(5.8), Inches(4.0))
txBox.word_wrap = True
tf = txBox.text_frame
tf.word_wrap = True

items = [
    ("Counselor Management", "Create, edit, and archive counselor accounts. Assign students to counselors with one click."),
    ("Student Roster", "Full CRUD on student profiles. View completion status and activity at a glance."),
    ("Access Control", "Developer → Admin → Counselor → Student hierarchy. Enforced at API and UI levels."),
    ("Data Integrity", "Counselors with active students cannot be deleted — protected by backend validation."),
]
first = True
for title, desc in items:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    run = p.add_run()
    run.text = f"▸  {title}"
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = CORAL
    run.font.name = BODY_FONT
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = f"    {desc}"
    run2.font.size = Pt(11)
    run2.font.color.rgb = CHARCOAL
    run2.font.name = BODY_FONT
    tf.add_paragraph()

add_rect(slide, Inches(6.8), Inches(1.1), Inches(6.1), Inches(5.9), fill_color=PALE_CORAL)
add_text_box(slide, "Key Capabilities at a Glance",
             Inches(7.0), Inches(1.25), Inches(5.7), Inches(0.5),
             font_size=13, bold=True, color=DARK_CORAL, font_name=HEADING_FONT)

stats = [
    ("4", "Permission levels\n(Developer, Admin,\nCounselor, Student)"),
    ("∞", "Students & counselors\nper admin account"),
    ("100%", "Role-enforced data\nisolation per firm"),
    ("Audit", "Complete activity trail\nfor compliance"),
]
for i, (num, label) in enumerate(stats):
    row = i // 2
    col = i % 2
    lft = Inches(7.1) + col * Inches(2.85)
    tp = Inches(2.0) + row * Inches(2.2)
    add_rect(slide, lft, tp, Inches(2.6), Inches(1.8), fill_color=WHITE,
             line_color=CORAL_LIGHT, line_width=Pt(0.5))
    add_text_box(slide, num, lft + Inches(0.1), tp + Inches(0.15),
                 Inches(2.4), Inches(0.7),
                 font_size=28, bold=True, color=CORAL, font_name=HEADING_FONT, align=PP_ALIGN.CENTER)
    add_text_box(slide, label, lft + Inches(0.1), tp + Inches(0.85),
                 Inches(2.4), Inches(0.8),
                 font_size=10, color=CHARCOAL, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────
# SLIDE 5: ADMIN — Customise the Platform
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Admin: Customise the Platform")
add_footer(slide, 5)

add_section_label(slide, "ADMIN  ·  2 of 3", Inches(0.4), Inches(1.05), Inches(2.2))

blocks = [
    ("Dynamic Profile Field Builder", CORAL, [
        "Create custom fields per student profile",
        "Field types: text, number, date, dropdown",
        "Mark fields as required or optional",
        "Organise into named categories",
        "Students only see fields from their admin",
    ]),
    ("University Repository", SLATE, [
        "Build an admin-scoped university list",
        "Add detailed info: rankings, programs, fees",
        "AI auto-extracts data from official sources",
        "Students only see universities you add",
        "Choose which data fields students can view",
    ]),
    ("Admin-Scoped Data Isolation", GREEN, [
        "Each firm's data is fully independent",
        "No cross-contamination between tenants",
        "Counselors inherit the admin's data scope",
        "Secure multi-tenancy by design",
    ]),
]

positions = [
    (Inches(0.4), Inches(1.45), Inches(4.0)),
    (Inches(4.65), Inches(1.45), Inches(4.0)),
    (Inches(8.9), Inches(1.45), Inches(4.0)),
]

for (title, color, bullets), (left, top, width) in zip(blocks, positions):
    add_rect(slide, left, top, width, Inches(5.4), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_rect(slide, left, top, width, Inches(0.45), fill_color=color)
    add_text_box(slide, title, left + Inches(0.15), top + Inches(0.05),
                 width - Inches(0.2), Inches(0.38),
                 font_size=12, bold=True, color=WHITE, font_name=HEADING_FONT)
    for j, b in enumerate(bullets):
        add_text_box(slide, f"▸  {b}",
                     left + Inches(0.2), top + Inches(0.65) + j * Inches(0.82),
                     width - Inches(0.35), Inches(0.75),
                     font_size=11, color=CHARCOAL)


# ─────────────────────────────────────────
# SLIDE 6: ADMIN — Intelligence & Oversight
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Admin: Intelligence & Oversight")
add_footer(slide, 6)

add_section_label(slide, "ADMIN  ·  3 of 3", Inches(0.4), Inches(1.05), Inches(2.2))

add_text_box(slide, "AI-Powered University Data Extraction",
             Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.55),
             font_size=16, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

steps = [
    ("1", "Trigger", "Admin adds a university to the repository"),
    ("2", "Multi-Source\nExtraction", "AI queries official website,\nCommon Data Set & web search"),
    ("3", "Store All Data", "Complete dataset saved to DB\nregardless of selection"),
    ("4", "Admin Curates", "Admin selects which fields\nstudents can see"),
    ("5", "Student View", "Students see only the\nadmin-approved fields"),
]

box_w = Inches(2.3)
box_h = Inches(1.4)
for i, (num, label, desc) in enumerate(steps):
    left = Inches(0.4) + i * (box_w + Inches(0.18))
    top = Inches(2.15)
    add_rect(slide, left, top, box_w, box_h, fill_color=PALE_CORAL,
             line_color=CORAL_LIGHT, line_width=Pt(0.5))
    add_text_box(slide, num, left + Inches(0.1), top + Inches(0.05),
                 Inches(0.35), Inches(0.35), font_size=14, bold=True, color=CORAL, font_name=HEADING_FONT)
    add_text_box(slide, label, left + Inches(0.1), top + Inches(0.32),
                 box_w - Inches(0.2), Inches(0.45),
                 font_size=11, bold=True, color=CHARCOAL)
    add_text_box(slide, desc, left + Inches(0.1), top + Inches(0.75),
                 box_w - Inches(0.2), Inches(0.6),
                 font_size=9, color=SLATE)
    if i < 4:
        add_text_box(slide, "→", left + box_w + Inches(0.02), top + Inches(0.5),
                     Inches(0.2), Inches(0.4),
                     font_size=14, color=CORAL, align=PP_ALIGN.CENTER)

left_features = [
    ("Email & SMTP Configuration", "Database-stored SMTP settings for password resets, task reminders, and event notifications. No hardcoded credentials."),
    ("Push Notification Management", "Real-time device push notifications for tasks, messages, activity feedback, and planner event reminders (1 day prior)."),
]
right_features = [
    ("Global Task Oversight", "View and track all tasks across the entire platform. Monitor counselor activity and student progress."),
    ("Counselor Activity Notifications", "Admins receive notifications when counselors assign tasks, upload documents, or schedule planner events."),
]

for i, (title, desc) in enumerate(left_features):
    top = Inches(3.75) + i * Inches(1.5)
    add_rect(slide, Inches(0.4), top, Inches(6.1), Inches(1.3), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_text_box(slide, f"▸  {title}", Inches(0.6), top + Inches(0.1), Inches(5.7), Inches(0.4),
                 font_size=12, bold=True, color=CHARCOAL)
    add_text_box(slide, desc, Inches(0.6), top + Inches(0.5), Inches(5.7), Inches(0.7),
                 font_size=10, color=SLATE)

for i, (title, desc) in enumerate(right_features):
    top = Inches(3.75) + i * Inches(1.5)
    add_rect(slide, Inches(6.9), top, Inches(6.1), Inches(1.3), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_text_box(slide, f"▸  {title}", Inches(7.1), top + Inches(0.1), Inches(5.7), Inches(0.4),
                 font_size=12, bold=True, color=CHARCOAL)
    add_text_box(slide, desc, Inches(7.1), top + Inches(0.5), Inches(5.7), Inches(0.7),
                 font_size=10, color=SLATE)


# ─────────────────────────────────────────
# SLIDE 7: COUNSELOR SECTION DIVIDER
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, CHARCOAL)

add_rect(slide, 0, 0, Inches(0.5), SLIDE_H, fill_color=SLATE)
add_rect(slide, Inches(0.5), 0, Inches(5.5), SLIDE_H, fill_color=RGBColor(0x4A, 0x5A, 0x6D))

circle3 = slide.shapes.add_shape(9, Inches(4.5), Inches(-1.0), Inches(5.5), Inches(5.5))
circle3.fill.solid()
circle3.fill.fore_color.rgb = RGBColor(0x3A, 0x4A, 0x5D)
circle3.line.fill.background()

add_text_box(slide, "02",
             Inches(0.8), Inches(1.5), Inches(2.0), Inches(2.0),
             font_size=96, bold=True, color=WHITE, font_name=HEADING_FONT)
add_text_box(slide, "COUNSELOR",
             Inches(0.8), Inches(3.6), Inches(4.5), Inches(0.8),
             font_size=30, bold=True, color=WHITE, font_name=HEADING_FONT)
add_rect(slide, Inches(0.8), Inches(4.5), Inches(1.5), Inches(0.06), fill_color=WHITE)
add_text_box(slide, "Focus on the students who\nneed you — guided by\ndata, tools, and clarity.",
             Inches(0.8), Inches(4.7), Inches(4.5), Inches(1.5),
             font_size=15, color=RGBColor(0xBB, 0xC4, 0xD0), font_name=HEADING_FONT, italic=True)

add_text_box(slide, "3 slides in this section:",
             Inches(6.5), Inches(2.2), Inches(6.5), Inches(0.4),
             font_size=12, color=SLATE, bold=True)
section_slides = [
    "My Caseload — scoped student access, messaging, planner",
    "Drive Student Progress — task workflow, activity feedback",
    "Documents & University Research — secure docs, uni data",
]
for i, s in enumerate(section_slides):
    add_rect(slide, Inches(6.5), Inches(2.75) + i * Inches(1.2), Inches(6.3), Inches(1.0),
             fill_color=RGBColor(0x3A, 0x3A, 0x3A), line_color=SLATE, line_width=Pt(0.5))
    add_text_box(slide, f"{i+1}.  {s}",
                 Inches(6.7), Inches(2.85) + i * Inches(1.2), Inches(5.9), Inches(0.8),
                 font_size=12, color=WHITE)

add_footer(slide, 7)


# ─────────────────────────────────────────
# SLIDE 8: COUNSELOR — My Caseload
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Counselor: My Caseload", accent_color=SLATE)
add_footer(slide, 8)

add_section_label(slide, "COUNSELOR  ·  1 of 3", Inches(0.4), Inches(1.05), Inches(2.5), color=SLATE)

add_text_box(slide, "Counselors see only what they need — their assigned students.",
             Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.5),
             font_size=15, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

features_left = [
    ("Scoped Student Access", "Counselors only see and interact with the students assigned to them by their admin. Full data isolation."),
    ("Rich Student Profiles", "View the student's full profile including custom admin-configured fields, completion status, and history."),
    ("Direct Messaging", "Built-in messaging panel for direct counselor-to-student communication, with notification alerts."),
]
features_right = [
    ("Planner Coordination", "View and add planner events for each student. Stay aligned on deadlines and appointments."),
    ("Notification Alerts", "Receive in-app and device push notifications for student activity, task completions, and messages."),
    ("Student Progress View", "At-a-glance dashboard of student task completion, profile percentage, and recent activities."),
]

for i, (title, desc) in enumerate(features_left):
    top = Inches(2.1) + i * Inches(1.5)
    add_rect(slide, Inches(0.4), top, Inches(6.0), Inches(1.3), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_rect(slide, Inches(0.4), top, Inches(0.12), Inches(1.3), fill_color=SLATE)
    add_text_box(slide, title, Inches(0.65), top + Inches(0.1), Inches(5.5), Inches(0.4),
                 font_size=12, bold=True, color=CHARCOAL)
    add_text_box(slide, desc, Inches(0.65), top + Inches(0.5), Inches(5.5), Inches(0.72),
                 font_size=10, color=SLATE)

for i, (title, desc) in enumerate(features_right):
    top = Inches(2.1) + i * Inches(1.5)
    add_rect(slide, Inches(6.85), top, Inches(6.0), Inches(1.3), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_rect(slide, Inches(6.85), top, Inches(0.12), Inches(1.3), fill_color=SLATE)
    add_text_box(slide, title, Inches(7.1), top + Inches(0.1), Inches(5.5), Inches(0.4),
                 font_size=12, bold=True, color=CHARCOAL)
    add_text_box(slide, desc, Inches(7.1), top + Inches(0.5), Inches(5.5), Inches(0.72),
                 font_size=10, color=SLATE)


# ─────────────────────────────────────────
# SLIDE 9: COUNSELOR — Drive Student Progress
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Counselor: Drive Student Progress", accent_color=SLATE)
add_footer(slide, 9)

add_section_label(slide, "COUNSELOR  ·  2 of 3", Inches(0.4), Inches(1.05), Inches(2.5), color=SLATE)

add_text_box(slide, "Task Assignment Workflow",
             Inches(0.4), Inches(1.5), Inches(8.0), Inches(0.45),
             font_size=14, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

task_steps = [
    ("Assign Task", "Counselor creates a task\nwith title, description\nand optional due date"),
    ("Auto Planner\nEvent", "System automatically\ncreates a planner entry\nfor due-date tasks"),
    ("Student Notified", "Student receives a\npush notification\nand in-app alert"),
    ("Track Progress", "Counselor monitors\nstatus: Pending →\nIn Progress → Done"),
]
for i, (label, desc) in enumerate(task_steps):
    left = Inches(0.4) + i * Inches(3.1)
    top = Inches(2.0)
    add_rect(slide, left, top, Inches(2.7), Inches(1.6), fill_color=PALE_CORAL,
             line_color=CORAL_LIGHT, line_width=Pt(0.5))
    add_text_box(slide, label, left + Inches(0.15), top + Inches(0.1), Inches(2.4), Inches(0.5),
                 font_size=12, bold=True, color=SLATE)
    add_text_box(slide, desc, left + Inches(0.15), top + Inches(0.6), Inches(2.4), Inches(0.9),
                 font_size=10, color=CHARCOAL)
    if i < 3:
        add_text_box(slide, "→", left + Inches(2.72), top + Inches(0.6), Inches(0.32), Inches(0.4),
                     font_size=14, color=SLATE, align=PP_ALIGN.CENTER)

add_text_box(slide, "Activity Feedback & Portfolio Review",
             Inches(0.4), Inches(3.8), Inches(8.0), Inches(0.45),
             font_size=14, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

feedback_items = [
    "Students log extracurricular activities with descriptions and outcomes",
    "Counselors review each activity entry and provide written feedback",
    "Feedback triggers a notification to the student instantly",
    "Activity history builds the student's portfolio for university applications",
    "Pagination keeps long activity lists manageable for both parties",
]
for i, item in enumerate(feedback_items):
    add_text_box(slide, f"▸  {item}", Inches(0.5), Inches(4.35) + i * Inches(0.48),
                 Inches(7.5), Inches(0.43), font_size=11, color=CHARCOAL)

add_rect(slide, Inches(8.9), Inches(1.1), Inches(4.1), Inches(5.9), fill_color=WHITE,
         line_color=LIGHT_GRAY, line_width=Pt(0.5))
add_rect(slide, Inches(8.9), Inches(1.1), Inches(4.1), Inches(0.45), fill_color=SLATE)
add_text_box(slide, "Planner Integration",
             Inches(9.05), Inches(1.12), Inches(3.7), Inches(0.4),
             font_size=12, bold=True, color=WHITE, font_name=HEADING_FONT)

planner_points = [
    "Full calendar view for\neach assigned student",
    "Tasks with due dates auto-\npopulate as planner events",
    "Counselor can add custom\nevents directly to the calendar",
    "Students get reminders\n1 day before events",
    "Scrollable event creation\nform for ease of use",
]
for i, point in enumerate(planner_points):
    add_text_box(slide, f"▸  {point}", Inches(9.1), Inches(1.7) + i * Inches(0.88),
                 Inches(3.7), Inches(0.8), font_size=11, color=CHARCOAL)


# ─────────────────────────────────────────
# SLIDE 10: COUNSELOR — Documents & University Research
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Counselor: Documents & University Research", accent_color=SLATE)
add_footer(slide, 10)

add_section_label(slide, "COUNSELOR  ·  3 of 3", Inches(0.4), Inches(1.05), Inches(2.5), color=SLATE)

add_rect(slide, Inches(0.4), Inches(1.45), Inches(6.0), Inches(5.6), fill_color=WHITE,
         line_color=LIGHT_GRAY, line_width=Pt(0.5))
add_rect(slide, Inches(0.4), Inches(1.45), Inches(6.0), Inches(0.45), fill_color=SLATE)
add_text_box(slide, "Secure Document Management",
             Inches(0.55), Inches(1.48), Inches(5.6), Inches(0.4),
             font_size=13, bold=True, color=WHITE, font_name=HEADING_FONT)

doc_items = [
    ("Upload Documents", "Counselors upload files (PDFs, Word docs, images) directly to each student's profile."),
    ("Visibility Control", "Toggle per-document whether the student can see it — useful for internal notes vs. shared resources."),
    ("Edit & Delete", "Full lifecycle management: rename, replace, or remove documents as the application process evolves."),
    ("Role-Based Access", "Only the assigned counselor and admin can see all documents. Student sees only what's shared with them."),
    ("Secure Storage", "Files stored with role-enforced access. Unauthorised access blocked at the API level."),
]
for i, (title, desc) in enumerate(doc_items):
    top = Inches(2.1) + i * Inches(0.95)
    add_text_box(slide, f"▸  {title}", Inches(0.6), top, Inches(5.5), Inches(0.3),
                 font_size=11, bold=True, color=SLATE)
    add_text_box(slide, desc, Inches(0.75), top + Inches(0.28), Inches(5.4), Inches(0.55),
                 font_size=10, color=CHARCOAL)

add_rect(slide, Inches(6.9), Inches(1.45), Inches(6.0), Inches(5.6), fill_color=WHITE,
         line_color=LIGHT_GRAY, line_width=Pt(0.5))
add_rect(slide, Inches(6.9), Inches(1.45), Inches(6.0), Inches(0.45), fill_color=SLATE)
add_text_box(slide, "University Research & Recommendations",
             Inches(7.05), Inches(1.48), Inches(5.6), Inches(0.4),
             font_size=13, bold=True, color=WHITE, font_name=HEADING_FONT)

uni_items = [
    ("Browse University Database", "Access the admin-scoped university list with all AI-extracted and admin-curated data fields."),
    ("Suggest Universities", "Recommend specific universities to students directly from the counselor portal."),
    ("View Rich Data", "See rankings, programs, acceptance rates, fees, and custom fields the admin has enabled."),
    ("Student Shortlisting", "Students can categorise universities as Reach, Target, or Safety for application planning."),
    ("Admin-Scoped Catalogue", "Counselors only see universities added by their admin — keeping the list relevant and focused."),
]
for i, (title, desc) in enumerate(uni_items):
    top = Inches(2.1) + i * Inches(0.95)
    add_text_box(slide, f"▸  {title}", Inches(7.1), top, Inches(5.5), Inches(0.3),
                 font_size=11, bold=True, color=SLATE)
    add_text_box(slide, desc, Inches(7.25), top + Inches(0.28), Inches(5.4), Inches(0.55),
                 font_size=10, color=CHARCOAL)


# ─────────────────────────────────────────
# SLIDE 11: STUDENT SECTION DIVIDER
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, CHARCOAL)

add_rect(slide, 0, 0, Inches(0.5), SLIDE_H, fill_color=BLUE)
add_rect(slide, Inches(0.5), 0, Inches(5.5), SLIDE_H, fill_color=DARK_BLUE)

circle4 = slide.shapes.add_shape(9, Inches(4.5), Inches(-1.0), Inches(5.5), Inches(5.5))
circle4.fill.solid()
circle4.fill.fore_color.rgb = RGBColor(0x14, 0x3D, 0x85)
circle4.line.fill.background()

add_text_box(slide, "03",
             Inches(0.8), Inches(1.5), Inches(2.0), Inches(2.0),
             font_size=96, bold=True, color=WHITE, font_name=HEADING_FONT)
add_text_box(slide, "STUDENT",
             Inches(0.8), Inches(3.6), Inches(4.5), Inches(0.8),
             font_size=36, bold=True, color=WHITE, font_name=HEADING_FONT)
add_rect(slide, Inches(0.8), Inches(4.5), Inches(1.5), Inches(0.06), fill_color=WHITE)
add_text_box(slide, "Own your journey — profile,\nportfolio, planner, and\nuniversity planning in one place.",
             Inches(0.8), Inches(4.7), Inches(4.5), Inches(1.5),
             font_size=15, color=PALE_BLUE, font_name=HEADING_FONT, italic=True)

add_text_box(slide, "3 slides in this section:",
             Inches(6.5), Inches(2.2), Inches(6.5), Inches(0.4),
             font_size=12, color=SLATE, bold=True)
section_slides = [
    "My Dashboard — tasks, notifications, profile completion",
    "Build My Portfolio — Activity Locker, Planner, Document Vault",
    "University Planning — shortlisting, Reach/Target/Safety",
]
for i, s in enumerate(section_slides):
    add_rect(slide, Inches(6.5), Inches(2.75) + i * Inches(1.2), Inches(6.3), Inches(1.0),
             fill_color=RGBColor(0x3A, 0x3A, 0x3A), line_color=BLUE, line_width=Pt(0.5))
    add_text_box(slide, f"{i+1}.  {s}",
                 Inches(6.7), Inches(2.85) + i * Inches(1.2), Inches(5.9), Inches(0.8),
                 font_size=12, color=WHITE)

add_footer(slide, 11)


# ─────────────────────────────────────────
# SLIDE 12: STUDENT — My Dashboard
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Student: My Dashboard", accent_color=BLUE)
add_footer(slide, 12)

add_section_label(slide, "STUDENT  ·  1 of 3", Inches(0.4), Inches(1.05), Inches(2.2), color=BLUE)

add_text_box(slide, "Everything a student needs, front and center.",
             Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.5),
             font_size=15, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

dashboard_widgets = [
    ("Profile Completion Tracker", BLUE, [
        "Visual progress indicator (% complete)",
        "Highlights which sections need attention",
        "Motivates students to keep profiles updated",
        "Instant feedback on completeness",
    ]),
    ("My Task List", CORAL, [
        "Tasks assigned by counselor or admin",
        "Status: Pending → In Progress → Completed",
        "Paginated list for easy navigation",
        "Due date displayed alongside each task",
    ]),
    ("Notifications & Reminders", GREEN, [
        "In-app notifications for new tasks & messages",
        "Device push notifications (with permission)",
        "Event reminders 1 day before planner events",
        "Custom reminder timing (minutes/hours/days)",
    ]),
]

for i, (title, color, bullets) in enumerate(dashboard_widgets):
    left = Inches(0.4) + i * Inches(4.3)
    top = Inches(2.05)
    add_rect(slide, left, top, Inches(4.0), Inches(4.8), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_rect(slide, left, top, Inches(4.0), Inches(0.45), fill_color=color)
    add_text_box(slide, title, left + Inches(0.15), top + Inches(0.06),
                 Inches(3.7), Inches(0.38),
                 font_size=12, bold=True, color=WHITE, font_name=HEADING_FONT)
    for j, b in enumerate(bullets):
        add_text_box(slide, f"▸  {b}", left + Inches(0.2), top + Inches(0.65) + j * Inches(0.92),
                     Inches(3.65), Inches(0.82), font_size=11, color=CHARCOAL)


# ─────────────────────────────────────────
# SLIDE 13: STUDENT — Build My Portfolio
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Student: Build My Portfolio", accent_color=BLUE)
add_footer(slide, 13)

add_section_label(slide, "STUDENT  ·  2 of 3", Inches(0.4), Inches(1.05), Inches(2.2), color=BLUE)

sections_portfolio = [
    ("Activity Locker", BLUE, [
        ("Log Activities", "Record extracurriculars,\nresponsibilities, volunteer\nwork, and achievements"),
        ("Describe Outcomes", "Detail what you learned\nor accomplished in\neach activity"),
        ("Counselor Feedback", "Receive written feedback\nfrom your counselor on\neach logged activity"),
        ("Build Over Time", "Activities accumulate to\nform a rich portfolio for\nuniversity applications"),
    ]),
    ("Personal Planner", SLATE, [
        ("Calendar View", "Full calendar interface\nwith month and week\nviews for planning"),
        ("Custom Events", "Create personal reminders\nfor deadlines, prep sessions,\nand important dates"),
        ("Smart Reminders", "Set reminders at custom\ntimes: minutes, hours,\nor days in advance"),
        ("Task Auto-Events", "Tasks with due dates from\nyour counselor appear\nautomatically in your planner"),
    ]),
    ("Document Vault", GREEN, [
        ("Shared Documents", "View documents your\ncounselor has shared\nwith you securely"),
        ("Application Files", "Store and access important\nfiles for university\napplication processes"),
        ("Role-Controlled", "See exactly what your\ncounselor has permitted —\nno more, no less"),
        ("Always Available", "Access your documents\nanytime via your secure\nstudent portal"),
    ]),
]

for i, (title, color, items) in enumerate(sections_portfolio):
    left = Inches(0.4) + i * Inches(4.3)
    top = Inches(1.45)
    add_rect(slide, left, top, Inches(4.0), Inches(5.7), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_rect(slide, left, top, Inches(4.0), Inches(0.45), fill_color=color)
    add_text_box(slide, title, left + Inches(0.15), top + Inches(0.06),
                 Inches(3.7), Inches(0.38),
                 font_size=13, bold=True, color=WHITE, font_name=HEADING_FONT)
    for j, (label, desc) in enumerate(items):
        item_top = top + Inches(0.62) + j * Inches(1.22)
        add_text_box(slide, label, left + Inches(0.2), item_top,
                     Inches(3.6), Inches(0.32),
                     font_size=11, bold=True, color=color)
        add_text_box(slide, desc, left + Inches(0.2), item_top + Inches(0.3),
                     Inches(3.6), Inches(0.82), font_size=10, color=CHARCOAL)


# ─────────────────────────────────────────
# SLIDE 14: STUDENT — University Planning
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Student: University Planning", accent_color=BLUE)
add_footer(slide, 14)

add_section_label(slide, "STUDENT  ·  3 of 3", Inches(0.4), Inches(1.05), Inches(2.2), color=BLUE)

add_text_box(slide, "Plan your university applications with guidance and data.",
             Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.5),
             font_size=15, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

categories = [
    ("Reach", CORAL, "Universities where admission\nis competitive but possible with\na strong application package."),
    ("Target", BLUE, "Universities where your profile\naligns well with typical admitted\nstudents — your primary focus."),
    ("Safety", GREEN, "Universities where you are\nconfidently within or above\ntheir typical admit profile."),
]

for i, (label, color, desc) in enumerate(categories):
    left = Inches(0.4) + i * Inches(4.3)
    top = Inches(2.1)
    add_rect(slide, left, top, Inches(4.0), Inches(2.3), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_rect(slide, left, top, Inches(4.0), Inches(0.5), fill_color=color)
    add_text_box(slide, label, left + Inches(0.15), top + Inches(0.07),
                 Inches(3.7), Inches(0.42),
                 font_size=14, bold=True, color=WHITE, font_name=HEADING_FONT)
    add_text_box(slide, desc, left + Inches(0.2), top + Inches(0.65),
                 Inches(3.65), Inches(1.5), font_size=11, color=CHARCOAL)

add_text_box(slide, "What students can see about each university:",
             Inches(0.4), Inches(4.55), Inches(12.5), Inches(0.4),
             font_size=13, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

data_fields = [
    "Official rankings and reputation data",
    "Acceptance rates and application deadlines",
    "Available programs and degree offerings",
    "Tuition fees and financial aid information",
    "Campus location and key facts",
    "Admin-curated notes and highlights",
]
for i, field in enumerate(data_fields):
    col = i % 2
    row = i // 2
    left = Inches(0.5) + col * Inches(6.4)
    top = Inches(5.1) + row * Inches(0.52)
    add_text_box(slide, f"▸  {field}", left, top, Inches(6.0), Inches(0.45),
                 font_size=11, color=CHARCOAL)

add_rect(slide, Inches(0.4), Inches(6.65), Inches(12.5), Inches(0.3), fill_color=PALE_BLUE)
add_text_box(slide, "  Students only see data fields and universities that their admin has approved — ensuring curated, relevant information.",
             Inches(0.6), Inches(6.67), Inches(12.0), Inches(0.28),
             font_size=9, color=DARK_BLUE)


# ─────────────────────────────────────────
# SLIDE 15: PLATFORM ARCHITECTURE
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Platform Architecture & Security")
add_footer(slide, 15)

add_rect(slide, Inches(0.4), Inches(1.1), Inches(5.9), Inches(6.0), fill_color=WHITE,
         line_color=LIGHT_GRAY, line_width=Pt(0.5))
add_rect(slide, Inches(0.4), Inches(1.1), Inches(5.9), Inches(0.45), fill_color=CHARCOAL)
add_text_box(slide, "Technology Stack",
             Inches(0.55), Inches(1.13), Inches(5.5), Inches(0.4),
             font_size=13, bold=True, color=WHITE, font_name=HEADING_FONT)

stack_layers = [
    ("Frontend", CORAL, [
        "React 18 + TypeScript",
        "Vite build tool",
        "Tailwind CSS + shadcn/ui (Radix UI)",
        "TanStack Query (server state)",
        "React Hook Form + Zod validation",
        "Wouter (client-side routing)",
    ]),
    ("Backend", SLATE, [
        "Node.js + Express.js (TypeScript, ES modules)",
        "Drizzle ORM",
        "bcrypt password hashing",
        "Express sessions + connect-pg-simple",
        "Nodemailer (email) + web-push",
    ]),
    ("Database", GREEN, [
        "PostgreSQL via Neon serverless",
        "Schema: Users, Students, Counselors,",
        "  Profile Fields, Universities, Tasks,",
        "  Activities, Planner Events",
    ]),
]

y = Inches(1.65)
for layer_name, color, items in stack_layers:
    add_rect(slide, Inches(0.55), y, Inches(1.1), Inches(0.3), fill_color=color)
    add_text_box(slide, layer_name, Inches(0.6), y + Inches(0.03), Inches(1.0), Inches(0.25),
                 font_size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for item in items:
        y += Inches(0.38)
        add_text_box(slide, f"  ·  {item}", Inches(1.75), y - Inches(0.02),
                     Inches(4.3), Inches(0.35), font_size=10, color=CHARCOAL)
    y += Inches(0.3)

add_rect(slide, Inches(6.7), Inches(1.1), Inches(6.2), Inches(6.0), fill_color=WHITE,
         line_color=LIGHT_GRAY, line_width=Pt(0.5))
add_rect(slide, Inches(6.7), Inches(1.1), Inches(6.2), Inches(0.45), fill_color=CHARCOAL)
add_text_box(slide, "Security & Multi-Tenancy",
             Inches(6.85), Inches(1.13), Inches(5.8), Inches(0.4),
             font_size=13, bold=True, color=WHITE, font_name=HEADING_FONT)

sec_items = [
    ("RBAC — 4 Permission Levels", CORAL,
     "Developer → Admin → Counselor → Student.\nEvery API endpoint and UI component enforces the hierarchy."),
    ("bcrypt Password Hashing", SLATE,
     "All passwords salted and hashed with bcrypt.\nNo plaintext credentials stored anywhere."),
    ("Session Management", GREEN,
     "Express sessions persisted in PostgreSQL via connect-pg-simple.\nSessions expire and rotate securely."),
    ("Multi-Tenancy", CORAL,
     "Each admin is an isolated tenant. Students, counselors, fields, and universities are scoped to their admin."),
    ("Data Integrity Guards", SLATE,
     "Counselors cannot be deleted while students are assigned. Validated at backend API and frontend UI."),
    ("Email/Password Auth", GREEN,
     "Custom email/password authentication with secure session tokens. No third-party OAuth dependency."),
]

for i, (title, color, desc) in enumerate(sec_items):
    top = Inches(1.7) + i * Inches(0.85)
    add_rect(slide, Inches(6.85), top, Inches(0.08), Inches(0.65), fill_color=color)
    add_text_box(slide, title, Inches(7.05), top + Inches(0.02), Inches(5.65), Inches(0.28),
                 font_size=11, bold=True, color=CHARCOAL)
    add_text_box(slide, desc, Inches(7.05), top + Inches(0.3), Inches(5.65), Inches(0.5),
                 font_size=9, color=SLATE)


# ─────────────────────────────────────────
# SLIDE 16: SUMMARY
# ─────────────────────────────────────────
slide = prs.slides.add_slide(blank_layout)
set_bg(slide, OFF_WHITE)
add_header_bar(slide, "Why Synergasia?")
add_footer(slide, 16)

add_text_box(slide, "One platform. Three empowered roles.",
             Inches(0.4), Inches(1.15), Inches(12.5), Inches(0.55),
             font_size=18, bold=True, color=CHARCOAL, font_name=HEADING_FONT, align=PP_ALIGN.CENTER)

persona_summaries = [
    ("Admin", CORAL, "Total control over your counseling firm — manage teams, customise the platform, and leverage AI to build a world-class university database with zero manual research."),
    ("Counselor", SLATE, "Focus entirely on your assigned students — assign tasks, review portfolios, share documents, and coordinate plans in one seamlessly scoped workspace."),
    ("Student", BLUE, "A personal command center for university applications — track tasks, log activities, plan with smart reminders, and explore curated university options with confidence."),
]

for i, (persona, color, summary) in enumerate(persona_summaries):
    left = Inches(0.4) + i * Inches(4.3)
    top = Inches(1.9)
    add_rect(slide, left, top, Inches(4.0), Inches(2.5), fill_color=WHITE,
             line_color=LIGHT_GRAY, line_width=Pt(0.5))
    add_rect(slide, left, top, Inches(4.0), Inches(0.5), fill_color=color)
    add_text_box(slide, persona, left + Inches(0.15), top + Inches(0.07),
                 Inches(3.7), Inches(0.42),
                 font_size=15, bold=True, color=WHITE, font_name=HEADING_FONT)
    add_text_box(slide, summary, left + Inches(0.2), top + Inches(0.65),
                 Inches(3.65), Inches(1.7), font_size=11, color=CHARCOAL)

add_text_box(slide, "3 Differentiators That Set Synergasia Apart",
             Inches(0.4), Inches(4.6), Inches(12.5), Inches(0.45),
             font_size=15, bold=True, color=CHARCOAL, font_name=HEADING_FONT)

differentiators = [
    ("AI-Powered Data Extraction",
     "Automatically pulls university information from multiple authoritative sources — no manual data entry. Admins curate what students see; the system captures everything."),
    ("Fully Dynamic Profiles",
     "Admins define exactly what information students need to provide. Custom fields, categories, and visibility rules mean the platform adapts to every firm's process."),
    ("Enterprise-Grade Multi-Tenant RBAC",
     "Four permission levels with data isolation between firms. Every action is authorised by role — at the API, at the query, and at the UI."),
]

for i, (title, desc) in enumerate(differentiators):
    left = Inches(0.4) + i * Inches(4.3)
    top = Inches(5.15)
    add_rect(slide, left, top, Inches(4.0), Inches(1.7), fill_color=PALE_CORAL,
             line_color=CORAL_LIGHT, line_width=Pt(0.5))
    add_text_box(slide, title, left + Inches(0.2), top + Inches(0.1),
                 Inches(3.65), Inches(0.4),
                 font_size=11, bold=True, color=DARK_CORAL, font_name=HEADING_FONT)
    add_text_box(slide, desc, left + Inches(0.2), top + Inches(0.5),
                 Inches(3.65), Inches(1.1), font_size=10, color=CHARCOAL)


# ─────────────────────────────────────────
# VERIFY AND SAVE
# ─────────────────────────────────────────
output_path = "synergasia_l0.pptx"
prs.save(output_path)
print(f"Saved: {output_path}")
print(f"Slide count: {len(prs.slides)}")
assert len(prs.slides) == TOTAL_SLIDES, f"Expected {TOTAL_SLIDES} slides, got {len(prs.slides)}"
print("Slide count check passed.")
